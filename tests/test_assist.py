"""Checks of the `assist` channel.

There are two classes of errors here, and both are silent. The first: the profile was saved
but did not reach the prompt; the model answers like a textbook, and you can tell only from
the quality of the answers. The second: the filter lets chatter through; the overlay floods
the screen with text exactly when the other person needs to be listened to.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from argus.app import Argus
from argus.assist import materials
from argus.assist.gate import AssistGate, is_small_talk
from argus.assist.prompt import LANGUAGE_RULES, compose, install
from argus.assist.settings import Answering, AssistSettings, read
from argus.brain.gate import GateSettings, ReplyMode
from argus.context.timeline import Entry, Kind
from argus.core.config import Config
from argus.core.events import Source
from argus.tts.settings import read as tts_read
from argus.ui.admin import apply_assist, build_panel, sections

NOW = 10**12


def app_with(**channel) -> Argus:
    cfg = Config()
    if channel:
        cfg.channel["assist"] = channel
    return Argus(cfg, with_screen=False, with_brain=False)


def said(text: str, source: Source = Source.THEM) -> Entry:
    return Entry(t=NOW, kind=Kind.SPEECH, source=source, text=text)


# --------------------------------------------------------------------------
# Profile -> prompt
# --------------------------------------------------------------------------


def test_profile_reaches_the_prompt():
    """The main item of the branch plan: the model answers for a specific person."""
    s = AssistSettings.model_validate({
        "profile": {"name": "Anton", "role": "Python developer", "years": 7,
                    "stack": "asyncio, PostgreSQL", "story": "left for the scale"},
    })
    text = compose(s)
    assert "Anton" in text and "asyncio" in text and "left for the scale" in text


def test_empty_profile_leaves_no_empty_headings():
    """An empty section in the prompt is noise that is paid for in tokens."""
    text = compose(AssistSettings())
    assert "# Context" not in text
    assert "# About this conversation" not in text


def test_prompt_is_derived_not_stored():
    """The prompt is built anew, not accumulated.

    If it were built on top of itself, every profile edit would double the system part, and this
    would be discovered by the token bill.
    """
    cfg = Config()
    cfg.channel["assist"] = {"profile": {"name": "Anton"}}
    first = install(cfg)
    second = install(cfg)
    assert first == second


def test_default_format_adds_nothing_extra():
    """The ordinary conversational format adds nothing.

    "Briefly" is already said in the core rules; a second line about the same thing in the
    prompt would be noise paid for in tokens on every question.
    """
    line = compose(AssistSettings.model_validate({"answering": {"format": "line"}}))
    default = compose(AssistSettings())
    assert line == default


def test_non_default_format_adds_its_own_hint():
    """The "code" and "long" formats — a difference that the core rules do not cover."""
    code = compose(AssistSettings.model_validate({"answering": {"format": "code"}}))
    long_ = compose(AssistSettings.model_validate({"answering": {"format": "long"}}))
    assert "Code first" in code
    assert "Three to five points" in long_


def test_language_is_the_conversation_until_it_is_chosen():
    """The default follows the conversation, not the saved value.

    "Same as conversation" is more accurate than a saved language: people
    switch to English in mid-phrase, and an explicit choice would only get in the way. The rule
    lives entirely in the channel settings, not in the core prompt, but by default it still
    sounds, it just does not argue with the conversation by naming a specific language.
    """
    text = compose(AssistSettings())
    assert "answer in the language of the conversation" in text.lower()
    # No rule that names a specific language may be in the prompt.
    for code, rule in LANGUAGE_RULES.items():
        if code:
            assert rule not in text, code


def test_chosen_language_overrides_the_conversation():
    """The point of making a choice is not having to switch along with each phrase."""
    text = compose(AssistSettings.model_validate({"profile": {"language": "en"}}))
    assert LANGUAGE_RULES["en"] in text and "Always answer in English" in text


def test_business_tone_is_the_default():
    """Business tone is the default — the phrase about profanity must not surface by itself."""
    text = compose(AssistSettings())
    assert "The tone is switched to rude" not in text


def test_rude_tone_adds_its_block():
    """Choosing the mode turns on a ready-made block as a whole, rather than assembling it from
    a field.
    """
    text = compose(AssistSettings.model_validate({"prompting": {"tone": "rude"}}))
    assert "The tone is switched to rude" in text


@pytest.mark.parametrize(
    "code,marker",
    [
        ("technical", "The tone is switched to technical"),
        ("neutral", "The tone is switched to ordinary"),
        ("friendly", "The tone is switched to friendly"),
        ("strict", "The tone is switched to strict"),
        ("trolling", "The tone is switched to trolling"),
    ],
)
def test_new_tones_add_their_own_block(code, marker):
    """Each new mode gets its own ready-made block and does not reuse another's."""
    text = compose(AssistSettings.model_validate({"prompting": {"tone": code}}))
    assert marker in text


# --------------------------------------------------------------------------
# Materials -> prompt
# --------------------------------------------------------------------------


def cfg_with_materials(tmp_path, files: dict[str, str], **materials) -> Config:
    """A config whose root is a temporary folder with lectures."""
    folder = tmp_path / "materials"
    folder.mkdir()
    for name, text in files.items():
        (folder / name).write_text(text, encoding="utf-8")
    cfg = Config()
    cfg.root = tmp_path
    cfg.channel["assist"] = {"materials": materials} if materials else {}
    return cfg


def test_lecture_reaches_the_prompt(tmp_path):
    """This is what materials exist for: the model answers from the course."""
    cfg = cfg_with_materials(tmp_path, {"01-indexes.md": "B-tree keeps the order"})
    text = install(cfg)
    assert "B-tree keeps the order" in text
    assert "# Materials" in text


def test_pdf_is_read_as_text(tmp_path):
    """Lectures are handed out as PDF, so we test on it, not only on a text file.

    The sample `tests/data/lecture.pdf` is made by hand and is small on purpose: a real lecture
    in the repository would be megabytes of someone else's material, and exactly one thing is
    tested here: that the text layer reaches the prompt.
    """
    cfg = cfg_with_materials(tmp_path, {})
    shutil.copy(Path(__file__).parent / "data" / "lecture.pdf",
                tmp_path / "materials" / "lecture.pdf")
    assert "B-tree keeps the order" in install(cfg)


def test_docx_keeps_the_tables(tmp_path):
    """In lectures on databases and networks, half of the content is in tables."""
    from docx import Document

    cfg = cfg_with_materials(tmp_path, {})
    d = Document()
    d.add_paragraph("B-tree keeps the order")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "hash"
    t.rows[0].cells[1].text = "O(1)"
    d.save(tmp_path / "materials" / "lecture.docx")

    text = install(cfg)
    assert "B-tree keeps the order" in text
    assert "hash | O(1)" in text


def test_pptx_takes_the_speaker_notes(tmp_path):
    """A slide has three words and a picture; what is said about them is in the notes."""
    from pptx import Presentation

    cfg = cfg_with_materials(tmp_path, {})
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[5])
    slide.shapes.title.text = "Transactions"
    slide.notes_slide.notes_text_frame.text = "ACID in full"
    p.save(tmp_path / "materials" / "slides.pptx")

    text = install(cfg)
    assert "Transactions" in text and "ACID in full" in text


def photo(path: Path, size=(1200, 800)) -> None:
    from PIL import Image

    Image.new("RGB", size, "white").save(path, quality=90)


def test_photo_goes_to_the_model_as_a_picture(tmp_path):
    """A photograph has no text layer, but it does not stop being a lecture."""
    cfg = cfg_with_materials(tmp_path, {})
    photo(tmp_path / "materials" / "page.jpg")
    m = materials.load(cfg)
    assert len(m.images) == 1
    assert not m.text
    assert m.items[0].pages == 1


def test_scanned_pdf_is_rendered_not_dropped(tmp_path):
    """A scan is a lecture without a text layer, not an empty file."""
    from PIL import Image

    cfg = cfg_with_materials(tmp_path, {})
    Image.new("RGB", (600, 800), "white").save(
        tmp_path / "materials" / "scan.pdf", "PDF"
    )
    m = materials.load(cfg)
    assert len(m.images) == 1
    assert not m.items[0].error


def test_rendered_pages_survive_a_restart(tmp_path):
    """Otherwise seven seconds of rendering are paid at every launch, not once."""
    from PIL import Image

    cfg = cfg_with_materials(tmp_path, {})
    scan = tmp_path / "materials" / "scan.pdf"
    Image.new("RGB", (600, 800), "white").save(scan, "PDF")

    first = materials.load(cfg).images
    materials._cache.clear()                      # as after a restart
    assert materials.load(cfg).images == first
    assert list((tmp_path / "materials" / materials.CACHE_DIRNAME).glob("*.jpg"))


def test_reworked_lecture_leaves_no_stale_pages(tmp_path):
    """Without cleanup the folder grows with every edit, and in a month there is a gigabyte in
    there.
    """
    from PIL import Image

    cfg = cfg_with_materials(tmp_path, {})
    scan = tmp_path / "materials" / "scan.pdf"
    page = Image.new("RGB", (600, 800), "white")
    page.save(scan, "PDF", save_all=True, append_images=[page, page])
    materials.load(cfg)

    page.save(scan, "PDF")                        # the lecture was rewritten as a single page
    materials._cache.clear()
    materials.load(cfg)
    kept = list((tmp_path / "materials" / materials.CACHE_DIRNAME).glob("*.jpg"))
    assert len(kept) == 1


def test_pages_have_their_own_ceiling(tmp_path):
    """A page as an image is fifteen hundred tokens against three hundred for text."""
    cfg = cfg_with_materials(tmp_path, {})
    for i in range(3):
        photo(tmp_path / "materials" / f"{i}.jpg")
    m = materials.load(cfg, max_pages=2)
    assert len(m.images) == 2
    assert m.dropped_pages == 1


def test_pages_reach_the_assistant(tmp_path):
    """A picture cannot go into the system part — it travels as message blocks."""
    from types import SimpleNamespace

    cfg = cfg_with_materials(tmp_path, {})
    photo(tmp_path / "materials" / "page.jpg")
    app = SimpleNamespace(assistant=SimpleNamespace(pinned=[], warm_soon=lambda: None))
    install(cfg, app)
    assert len(app.assistant.pinned) == 1


def test_no_materials_no_rule(tmp_path):
    """The rule "answer from the materials" without materials is an order to search for
    emptiness.
    """
    text = install(cfg_with_materials(tmp_path, {}))
    assert "# Materials" not in text
    assert "Answer from them" not in text


def test_materials_can_be_switched_off_without_deleting(tmp_path):
    """Between exams a lecture is not needed and costs tokens on every request."""
    cfg = cfg_with_materials(tmp_path, {"lecture.txt": "lecture text"}, enabled=False)
    assert "lecture text" not in install(cfg)


def test_materials_keep_the_order_of_names(tmp_path):
    """Lectures are numbered. The order "which file was touched last" is not an order."""
    m = materials.load(cfg_with_materials(
        tmp_path, {"02-second.txt": "second", "01-first.txt": "first"}
    ))
    assert m.text.index("first") < m.text.index("second")


def test_oversized_materials_are_cut_and_counted(tmp_path):
    """To cut silently is to lie about what the model answers from."""
    m = materials.load(cfg_with_materials(tmp_path, {"course.txt": "a" * 500}), max_chars=100)
    assert len(m.text) == 100
    assert m.dropped > 0


def test_scan_without_text_is_named_not_skipped(tmp_path):
    """An empty file in the list without an explanation — an exam with no lecture, which nobody
    knew about.
    """
    m = materials.load(cfg_with_materials(tmp_path, {"scan.txt": "   "}))
    assert m.items[0].error
    assert not m.text


def test_broken_channel_settings_do_not_break_startup():
    """You cannot die at startup because of a setting that the app itself wrote."""
    cfg = Config()
    cfg.channel["assist"] = {"answering": {"min_words": "many"}}
    assert read(cfg).answering.min_words == Answering().min_words


# --------------------------------------------------------------------------
# Filter "question or chatter"
# --------------------------------------------------------------------------

# Chatter and "too few words" are a prefilter of `AssistGate`, which is now applied only to
# `ReplyMode.SEMANTIC` (see `assist/gate.py`): the default `BY_FORM` does not pass through it
# at all. So all tests in this section set `system_mode=SEMANTIC` explicitly, otherwise they
# would be testing the wrong filter.
#
# The inputs of the gate tests below are Russian on purpose: the stop phrases (`Answering.stop_phrases`)
# and the question and task words of the core gate (`brain/gate.py`) are Russian word lists.
def gate(**answering) -> AssistGate:
    return AssistGate(
        settings=GateSettings(cooldown_s=0, system_mode=ReplyMode.SEMANTIC),
        answering=Answering(**answering),
    )


@pytest.mark.parametrize("text", [
    "Слышно меня?",
    "Приятно познакомиться!",
    "Ага, понял, спасибо",
    "Хорошо, понял.",
])
def test_small_talk_is_ignored(text):
    assert not gate().on_speech(said(text), NOW).wake


@pytest.mark.parametrize("text", [
    "Расскажите про устройство GIL в CPython",
    "Как бы вы шардировали таблицу на сто миллионов строк?",
    "Напишите функцию, которая разворачивает связный список",
])
def test_real_questions_pass(text):
    assert gate().on_speech(said(text), NOW).wake


def test_by_form_answers_even_trivial_questions():
    """`BY_FORM`, the default for system sound, does not know the concept of "chatter".

    "Слышно меня?" (Can you hear me?) is grammatically a question, and by form that is enough:
    the mode deliberately does not check whether it is a "real" question by meaning. That is the
    price at which it is free and instant. Those who want to filter such lines switch to
    `SEMANTIC` (see `gate()` above).
    """
    by_form = AssistGate(settings=GateSettings(cooldown_s=0, system_mode=ReplyMode.BY_FORM))
    assert by_form.on_speech(said("Слышно меня?"), NOW).wake
    # Not a question at all, neither by form nor by meaning, so it is still silent.
    assert not by_form.on_speech(said("Ага, понял, спасибо"), NOW).wake


def test_stop_phrase_inside_a_real_question_does_not_kill_it():
    """A stop phrase silences a line only when it takes up almost all of it.

    Substring search without this correction would silently swallow real questions, and it
    would look like a recognition breakage, not like the filter working.
    """
    assert not is_small_talk(
        "А как вы поняли, что всё хорошо и можно катить в прод?",
        Answering().stop_phrases,
    )


def test_words_not_characters():
    """The core threshold in characters is wrong in both directions at once."""
    short_question = "Что такое GIL?"     # 14 characters, 3 words
    polite = "Ага, понял, спасибо"        # 19 characters, 3 words
    g = gate(min_words=3)
    assert g.on_speech(said(short_question), NOW).wake
    assert not g.on_speech(said(polite), NOW).wake


def test_own_speech_is_never_answered():
    assert not gate().on_speech(said("What is a GIL?", Source.ME), NOW).wake


def test_counters_are_visible():
    """Without counters, filter refusals look like a breakage."""
    g = gate()
    g.on_speech(said("Слышно меня?"), NOW)
    g.on_speech(said("Ok"), NOW)
    stats = g.stats()
    assert stats["chatter"] == 1 and stats["too short"] == 1
    assert stats["filtered out"] == 2


# --------------------------------------------------------------------------
# Panel
# --------------------------------------------------------------------------


def test_panel_installs_prompt_and_gate_at_startup():
    """Without installation at startup, the first session would run on the core defaults."""
    app = app_with(profile={"name": "Anton"})
    app.assistant = None
    build_panel(app)
    assert "Anton" in app.cfg.brain.system_prompt


def test_editing_profile_rebuilds_the_prompt(tmp_path):
    app = app_with()
    app.assistant = None
    # A bare `Config()` has the current folder as its root, which is the working copy: without
    # this line the edit below would be written to the developer's real `argus.local.json`.
    app.cfg.root = tmp_path
    panel = build_panel(app)
    panel.apply({"channel.assist.profile.summary": "Anton, built an event bus"})
    assert "Anton" in app.cfg.brain.system_prompt


# The answer format has been removed from the page, and `panel.apply` no longer lets it
# through: an edit is filtered by the declared fields. The applier itself is alive: on reset
# `changed` comes from the file, not from the page, so we test it directly, not through the
# panel.


def test_editing_answering_rebuilds_the_gate():
    """An edit to the filter must reach the pipeline, not just the file."""
    cfg = Config()
    app = Argus(cfg, with_screen=False, with_brain=True)
    build_panel(app)
    cfg.channel["assist"] = {"answering": {"format": "code"}}
    apply_assist(app, cfg, {"channel.assist.answering.format"})
    assert app.assistant.gate.answering.format == "code"


def test_editing_mic_mode_reaches_the_gate_settings():
    """mic_mode lives in Answering, but it is read by the core's GateSettings:

    install() must carry the value from one to the other, otherwise the choice on the page is
    saved but affects nothing.
    """
    cfg = Config()
    app = Argus(cfg, with_screen=False, with_brain=True)
    build_panel(app)
    assert app.assistant.gate.settings.mic_mode is ReplyMode.OFF

    cfg.channel["assist"] = {"answering": {"mic_mode": "wake_word", "wake_word": "Argus"}}
    apply_assist(
        app, cfg,
        {"channel.assist.answering.mic_mode", "channel.assist.answering.wake_word"},
    )
    assert app.assistant.gate.settings.mic_mode is ReplyMode.WAKE_WORD
    assert app.assistant.gate.settings.wake_word == "Argus"


def test_editing_system_mode_reaches_the_gate_settings():
    """The same install() handoff as for mic_mode, but for the other person.

    The default is `BY_FORM`, not `OFF`: answering the other person was always supported, with
    no switch.
    """
    cfg = Config()
    app = Argus(cfg, with_screen=False, with_brain=True)
    build_panel(app)
    assert app.assistant.gate.settings.system_mode is ReplyMode.BY_FORM

    cfg.channel["assist"] = {"answering": {"system_mode": "wake_word"}}
    apply_assist(app, cfg, {"channel.assist.answering.system_mode"})
    assert app.assistant.gate.settings.system_mode is ReplyMode.WAKE_WORD


def test_gate_counters_survive_an_edit():
    """The counters are about the session, not about a setting: there is no need to wipe them.

    The chatter filter is now visible only on `SEMANTIC` (the default `BY_FORM` does not pass
    through it at all), so we switch `SEMANTIC` on explicitly, otherwise "Слышно меня?" (Can you
    hear me?) would just wake the model instead of settling into the counter.
    """
    app = Argus(Config(), with_screen=False, with_brain=True)
    build_panel(app)
    app.cfg.channel["assist"] = {"answering": {"system_mode": "semantic"}}
    apply_assist(app, app.cfg, {"channel.assist.answering.system_mode"})
    app.assistant.gate.on_speech(said("Слышно меня?"), NOW)
    app.cfg.channel["assist"] = {"answering": {"system_mode": "semantic", "format": "long"}}
    apply_assist(app, app.cfg, {"channel.assist.answering.format"})
    assert app.assistant.gate.stats()["chatter"] == 1


def test_every_setting_on_the_page_is_real():
    """A typo in a path gives a field that silently configures nothing."""
    app = app_with()
    app.assistant = None
    values = build_panel(app).values()
    for section in sections(tts_read(app.cfg)):
        for setting in section.settings:
            # The chat name is a property of the chat itself, not a config field: it appears in the values
            # only while the chat is open (`tests/test_chats.py`).
            if setting.path == "chat.name":
                continue
            assert setting.path in values, setting.path
