"""Materials: encodings, ceilings, page cache, lock.

Silent breakages live here: a file in a foreign encoding turned into noise without a single
error, a scan of six hundred pages was rendered whole for the sake of forty, photos were
re-encoded on every settings save, and a lecture over the server's ceiling got a refusal
instead of an answer.
"""

from __future__ import annotations

import codecs
import threading
import time
from pathlib import Path

from argus.assist import materials
from argus.assist.prompt import install
from argus.core.config import Config


def cfg_for(tmp_path: Path, provider: str = "claude-code") -> Config:
    (tmp_path / "materials").mkdir(exist_ok=True)
    cfg = Config()
    cfg.root = tmp_path
    cfg.brain.provider = provider
    return cfg


def folder(tmp_path: Path) -> Path:
    return tmp_path / "materials"


def photo(path: Path, size=(1200, 800)) -> None:
    from PIL import Image

    Image.new("RGB", size, "white").save(path, quality=90)


def scan(path: Path, pages: int) -> None:
    from PIL import Image

    page = Image.new("RGB", (600, 800), "white")
    page.save(path, "PDF", save_all=True, append_images=[page] * (pages - 1))


def make_pdf(texts: list[str]) -> bytes:
    """A real PDF with a text layer: one page per line."""
    n = len(texts)
    font_id = 3 + 2 * n
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", f"<< /Type /Pages /Kids [{kids}] /Count {n} >>"]
    for i, text in enumerate(texts):
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 400] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>")
        stream = f"BT /F1 12 Tf 20 350 Td ({text}) Tj ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{off:010d} 00000 n \n" for off in offsets).encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()
    return out


# --------------------------------------------------------------------------
# Encodings
# --------------------------------------------------------------------------

# Russian on purpose: the subject here is detecting the Windows-1251 encoding of Russian text.
NOTE = "Привет, это конспект лекции по базам данных. Индекс ускоряет поиск."


def test_a_windows_1251_note_is_read_as_russian_not_as_noise(tmp_path):
    """The old Notepad saves Russian text in "ANSI": it used to turn into replacement characters."""
    cfg = cfg_for(tmp_path)
    (folder(tmp_path) / "notes.txt").write_bytes(NOTE.encode("cp1251"))
    m = materials.load(cfg)
    assert NOTE in m.text
    assert "�" not in m.text
    assert not m.items[0].error


def test_a_utf_16_note_is_read(tmp_path):
    cfg = cfg_for(tmp_path)
    (folder(tmp_path) / "notes.txt").write_bytes(NOTE.encode("utf-16"))        # with a byte order mark
    m = materials.load(cfg)
    assert NOTE in m.text and "\x00" not in m.text


def test_a_utf_8_note_with_a_mark_is_clean(tmp_path):
    cfg = cfg_for(tmp_path)
    (folder(tmp_path) / "notes.md").write_bytes(codecs.BOM_UTF8 + NOTE.encode("utf-8"))
    assert materials.load(cfg).text.endswith(NOTE)


def test_utf_8_with_a_stray_byte_still_arrives_almost_whole(tmp_path):
    """Better half of the notes than nothing: one bad letter does not send the file to cp1251."""
    cfg = cfg_for(tmp_path)
    raw = (NOTE * 10).encode("utf-8") + b"\xff" + NOTE.encode("utf-8")
    (folder(tmp_path) / "notes.txt").write_bytes(raw)
    m = materials.load(cfg)
    assert NOTE in m.text
    assert "РџСЂ" not in m.text                                  # not mojibake from Windows-1251 (Russian on purpose)
    assert not m.items[0].error


def test_a_file_that_is_noise_is_an_error_in_the_list_not_n_characters(tmp_path):
    cfg = cfg_for(tmp_path)
    (folder(tmp_path) / "junk.txt").write_bytes(bytes(range(256)) * 20)
    m = materials.load(cfg)
    assert m.items[0].error
    assert not m.text
    assert "chars" not in materials.describe(m)


def test_a_latin_note_with_accents_is_not_turned_into_cyrillic(tmp_path, monkeypatch):
    """French text in "ANSI" is Latin with a couple of accented characters: a Cyrillic code page
    would turn "é" into a Cyrillic letter, and the list would show "N chars" without
    a shadow of doubt."""
    import locale

    cfg = cfg_for(tmp_path)
    french = "Résumé de la réunion : café et thé."
    (folder(tmp_path) / "notes.txt").write_bytes(french.encode("cp1252"))
    monkeypatch.setattr(locale, "getpreferredencoding", lambda do_setlocale=True: "UTF-8")
    m = materials.load(cfg)
    assert french in m.text and not m.items[0].error


def test_the_ansi_page_of_the_computer_reads_a_polish_note(tmp_path, monkeypatch):
    """On a Polish Windows, "ANSI" is 1250: "ś" and "ż" are not where they are in Western 1252."""
    import locale

    cfg = cfg_for(tmp_path)
    polish = "Wczoraj omawialiśmy zależności między tabelami: klucze główne, obce i ograniczenia."
    (folder(tmp_path) / "notatki.txt").write_bytes(polish.encode("cp1250"))
    monkeypatch.setattr(locale, "getpreferredencoding", lambda do_setlocale=True: "cp1250")
    assert polish in materials.load(cfg).text


def test_an_english_note_with_russian_words_stays_readable_in_both_languages(tmp_path):
    """Fewer than a quarter of the letters are Russian, but there are whole words: this is not
    Latin with diacritics.
    """
    cfg = cfg_for(tmp_path)
    note = "Database indexes speed up search. " * 3 + "Индексы ускоряют поиск."     # Russian on purpose: a mixed-language file
    (folder(tmp_path) / "mixed.txt").write_bytes(note.encode("cp1251"))
    assert note in materials.load(cfg).text


def test_a_russian_note_full_of_code_stays_russian(tmp_path):
    """There are many Latin words, but more than a quarter of the letters are Russian: this is
    not Latin with diacritics.
    """
    cfg = cfg_for(tmp_path)
    note = "SELECT name FROM users; -- выборка имён пользователей из таблицы"     # Russian on purpose
    (folder(tmp_path) / "sql.txt").write_bytes(note.encode("cp1251"))
    assert note in materials.load(cfg).text


# --------------------------------------------------------------------------
# Order in docx
# --------------------------------------------------------------------------


def test_a_docx_keeps_tables_where_the_author_put_them(tmp_path):
    """Previously all tables went to the end after all the paragraphs — "in their order" was not
    satisfied.
    """
    from docx import Document

    cfg = cfg_for(tmp_path)
    d = Document()
    d.add_paragraph("First an introductory paragraph")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "hash"
    t.rows[0].cells[1].text = "O(1)"
    d.add_paragraph("Then a conclusion about the table")
    d.save(folder(tmp_path) / "lecture.docx")

    text = materials.load(cfg).text
    assert text.index("First an") < text.index("hash | O(1)") < text.index("Then a conclusion")


# --------------------------------------------------------------------------
# Pages: ceiling, cache, photos
# --------------------------------------------------------------------------


def spy(monkeypatch, name: str) -> list:
    calls = []
    real = getattr(materials, name)

    def wrapper(*args, **kwargs):
        calls.append(args)
        return real(*args, **kwargs)

    monkeypatch.setattr(materials, name, wrapper)
    return calls


def test_rendering_stops_at_the_page_ceiling(tmp_path, monkeypatch):
    """A six-hundred-page scan must not be rendered whole for the sake of forty."""
    cfg = cfg_for(tmp_path)
    scan(folder(tmp_path) / "scan.pdf", 5)
    renders = spy(monkeypatch, "_render_pdf")

    m = materials.load(cfg, max_pages=2)

    assert len(m.images) == 2 and m.dropped_pages == 3
    assert [limit for _, limit in renders] == [2]
    on_disk = sorted(p.name for p in (folder(tmp_path) / materials.CACHE_DIRNAME).glob("*.jpg"))
    assert len(on_disk) == 2 and all("-l2-n5-" in n for n in on_disk)


def test_a_cache_cut_for_a_small_ceiling_is_not_given_to_a_bigger_one(tmp_path):
    """The page ceiling was raised — the earlier cut-off of two sheets must not be passed off as
    the lecture.
    """
    cfg = cfg_for(tmp_path)
    scan(folder(tmp_path) / "scan.pdf", 5)
    assert len(materials.load(cfg, max_pages=2).images) == 2

    assert len(materials.load(cfg, max_pages=4).images) == 4          # from memory
    materials._cache.clear()                                          # as after a restart
    assert len(materials.load(cfg, max_pages=5).images) == 5          # from disk
    assert materials.load(cfg, max_pages=5).dropped_pages == 0


def test_a_cache_cut_for_a_big_ceiling_serves_a_smaller_one_without_drawing_again(
        tmp_path, monkeypatch):
    cfg = cfg_for(tmp_path)
    scan(folder(tmp_path) / "scan.pdf", 5)
    first = materials.load(cfg, max_pages=4).images
    materials._cache.clear()
    renders = spy(monkeypatch, "_render_pdf")

    got = materials.load(cfg, max_pages=2)

    assert got.images == first[:2] and got.dropped_pages == 3
    assert renders == []


def test_a_photo_is_decoded_once_not_on_every_load(tmp_path, monkeypatch):
    """Previously every settings save re-decoded all photos: seconds in the loop."""
    cfg = cfg_for(tmp_path)
    photo(folder(tmp_path) / "page.jpg")
    opened = spy(monkeypatch, "_open_image")

    first = materials.load(cfg).images
    for _ in range(3):
        assert materials.load(cfg).images == first
    materials._cache.clear()                                          # restart: disk
    assert materials.load(cfg).images == first

    assert len(opened) == 1


def test_photos_over_the_ceiling_are_not_even_decoded(tmp_path, monkeypatch):
    cfg = cfg_for(tmp_path)
    for i in range(5):
        photo(folder(tmp_path) / f"{i}.jpg")
    opened = spy(monkeypatch, "_open_image")

    m = materials.load(cfg, max_pages=2)

    assert len(m.images) == 2 and m.dropped_pages == 3
    assert len(opened) == 2
    # A file that did not fit is named in the list, not left standing empty.
    assert [bool(i.error) for i in m.items] == [False, False, True, True, True]


def test_a_deleted_file_is_forgotten_by_the_memory_cache(tmp_path):
    cfg = cfg_for(tmp_path)
    (folder(tmp_path) / "a.txt").write_text("first", encoding="utf-8")
    (folder(tmp_path) / "b.txt").write_text("second", encoding="utf-8")
    materials.load(cfg)
    assert any(k[0].name == "b.txt" for k in materials._cache)

    (folder(tmp_path) / "b.txt").unlink()
    materials.load(cfg)
    assert not any(k[0].name == "b.txt" for k in materials._cache)


def test_a_thin_text_layer_does_not_hide_the_pages(tmp_path):
    """A watermark, page numbers or poor recognition — this is not a lecture in letters."""
    cfg = cfg_for(tmp_path)
    (folder(tmp_path) / "scan.pdf").write_bytes(make_pdf(["No. 1", "No. 2", "No. 3"]))

    m = materials.load(cfg)

    assert len(m.images) == 3
    assert "No. 1" in m.text                                          # the layer travels too
    assert m.items[0].pages == 3 and m.items[0].chars > 0
    assert "3 pages as images" in materials.describe(m)


def test_slides_with_a_couple_of_lines_stay_text_not_forty_pictures(tmp_path):
    """The sparseness threshold is low on purpose: slides with a couple of lines are text, and
    as pictures they would cost many times more on every question."""
    cfg = cfg_for(tmp_path)
    line = "B-tree keeps the keys sorted so the search is fast."                      # 51 characters
    (folder(tmp_path) / "slides.pdf").write_bytes(make_pdf([line] * 6))

    m = materials.load(cfg)

    assert m.images == [] and "keeps the keys sorted" in m.text


def test_a_real_text_layer_still_goes_as_text_only(tmp_path):
    cfg = cfg_for(tmp_path)
    page = "Index keeps the order of keys and speeds up the search. " * 4        # about 220 characters
    (folder(tmp_path) / "lecture.pdf").write_bytes(make_pdf([page, page]))

    m = materials.load(cfg)

    assert m.images == [] and "speeds up the search" in m.text
    assert m.items[0].pages == 0


# --------------------------------------------------------------------------
# Lock
# --------------------------------------------------------------------------


def test_two_loads_at_once_do_not_run_together(tmp_path, monkeypatch):
    """Cleaning the page cache of one parse would erase the fresh pages of another."""
    cfg = cfg_for(tmp_path)
    (folder(tmp_path) / "a.txt").write_text("text", encoding="utf-8")
    running = peak = 0
    real = materials._content_of

    def slow(*args, **kwargs):
        nonlocal running, peak
        running += 1
        peak = max(peak, running)
        time.sleep(0.1)
        running -= 1
        return real(*args, **kwargs)

    monkeypatch.setattr(materials, "_content_of", slow)
    threads = [threading.Thread(target=materials.load, args=(cfg,)) for _ in range(4)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert peak == 1


def test_install_inside_a_load_does_not_deadlock(tmp_path):
    """`install` calls `load` — the same thread enters under the lock again."""
    cfg = cfg_for(tmp_path)
    (folder(tmp_path) / "a.txt").write_text("lecture text", encoding="utf-8")
    with materials._LOCK:
        assert "lecture text" in install(cfg)


# --------------------------------------------------------------------------
# Account server ceilings
# --------------------------------------------------------------------------


def test_the_relay_ceiling_on_text_is_applied_and_counted(tmp_path):
    from argus.brain import relay_api

    cfg = cfg_for(tmp_path, "relay")
    (folder(tmp_path) / "course.txt").write_text("a" * (relay_api.MAX_SYSTEM_CHARS + 1000),
                                               encoding="utf-8")
    m = materials.load(cfg)
    cap = relay_api.MAX_SYSTEM_CHARS - materials.RELAY_SYSTEM_RESERVE
    assert len(m.text) == cap
    assert m.dropped > 0
    # Together with the rest of the prompt, the system part does not go over the server's ceiling.
    assert len(install(cfg)) < relay_api.MAX_SYSTEM_CHARS


def test_the_relay_ceiling_on_pages_leaves_a_place_for_the_screen(tmp_path):
    from argus.brain import relay_api

    cfg = cfg_for(tmp_path, "relay")
    for i in range(relay_api.MAX_IMAGES + 3):
        photo(folder(tmp_path) / f"{i:03d}.jpg", size=(60, 40))
    m = materials.load(cfg)               # the materials ceiling itself is not set, and it is cut anyway
    assert len(m.images) == relay_api.MAX_IMAGES - 1
    assert m.dropped_pages == 4


def test_the_relay_ceiling_on_page_size_is_applied(tmp_path, monkeypatch):
    from argus.brain import relay_api

    cfg = cfg_for(tmp_path, "relay")
    for i in range(12):
        photo(folder(tmp_path) / f"{i:02d}.jpg", size=(60, 40))
    monkeypatch.setattr(materials, "_jpeg", lambda img: b"x" * 1_500_000)         # about 2 MB of base64
    m = materials.load(cfg)
    sizes = sum((len(p) + 2) // 3 * 4 for p in m.images)
    assert 0 < len(m.images) < 12
    assert sizes <= relay_api.MAX_IMAGES_B64 - relay_api.MAX_IMAGE_B64
    assert m.dropped_pages == 12 - len(m.images)


def test_other_providers_are_not_cut_by_the_relays_ceilings(tmp_path):
    from argus.brain import relay_api

    cfg = cfg_for(tmp_path, "claude-code")
    (folder(tmp_path) / "course.txt").write_text("a" * (relay_api.MAX_SYSTEM_CHARS + 1000),
                                               encoding="utf-8")
    assert len(materials.load(cfg).text) > relay_api.MAX_SYSTEM_CHARS
